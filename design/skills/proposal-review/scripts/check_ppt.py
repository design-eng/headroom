#!/usr/bin/env python3
"""제안서 PPT 규격 검사기 — ppt-design-system-v2 기준.

읽기 전용. 파일을 수정하지 않고 위반 사항만 보고한다.
사용법:
    python3 check_ppt.py <파일.pptx> [--slides 29,36,59] [--json]
"""
import argparse, json, re, sys
from collections import defaultdict

try:
    from pptx import Presentation
except ImportError:
    sys.exit("python-pptx가 필요합니다:  pip install python-pptx")

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
EMU_PT = 12700

# ── ppt-design-system-v2 토큰 ────────────────────────────────────────────
PALETTE = {
    "1B2B41": "네이비", "3A88B4": "브랜드 블루", "66B2E0": "라이트블루",
    "D5D2CA": "웰그레이", "333333": "본문 블랙",
    "FFFFFF": "흰색", "000000": "검정",
}
# 구버전 색 → 치환 대상
REMAP = {
    **{c: "1B2B41" for c in ("002535", "122C43", "376092", "262626", "404040")},
    **{c: "3A88B4" for c in ("0085AD", "2F6FB7", "00629F", "00A650", "FFC000", "C00000")},
    **{c: "66B2E0" for c in ("4BBBE6", "50CFFF", "9DC3E6", "BDD7EE")},
    **{c: "D5D2CA" for c in ("BABABA", "CBCBCB", "D9D8D6", "E7E6E6")},
}
FONT_OK = re.compile(r"^Pretendard(\s+(SemiBold|Thin|Light|Medium|Bold|ExtraBold|Black))?$")
MARGIN = dict(left=46.0, right=914.0, top=99.0, bottom=464.0)
SLIDE_W, SLIDE_H = 960.0, 540.0
MIN_PT = 11.0
ROUND_EMU = 50800  # 4pt
TOL = 1.0          # 좌표 허용오차(pt) — 99.0812… 같은 값을 정수 비교하지 않는다


def pt(emu):
    return None if emu is None else round(emu / EMU_PT, 2)


def walk(shapes, path=""):
    """그룹 안까지 모든 도형을 (이름경로, shape)로 펼친다."""
    for sh in shapes:
        label = f"{path}{sh.shape_type and sh.name or sh.name}"
        yield label, sh
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # GROUP
            yield from walk(sh.shapes, path=f"{label} / ")


class Finding(dict):
    def __init__(self, slide, rule, shape, detail, fix=""):
        super().__init__(slide=slide, rule=rule, shape=shape, detail=detail, fix=fix)


def check_colors(el, slide, name, out):
    for clr in el.iter(f"{A}srgbClr"):
        val = (clr.get("val") or "").upper()
        if not val or val in PALETTE:
            continue
        fix = f"→ #{REMAP[val]} ({PALETTE[REMAP[val]]})" if val in REMAP else "팔레트 5색 중 하나로 치환"
        out.append(Finding(slide, "색상", name, f"#{val} 은 팔레트 외 색", fix))


def check_fonts(el, slide, name, out):
    for rpr in list(el.iter(f"{A}rPr")) + list(el.iter(f"{A}defRPr")) + list(el.iter(f"{A}endParaRPr")):
        faces = {}
        for tag in ("latin", "ea", "cs"):
            node = rpr.find(f"{A}{tag}")
            if node is not None:
                faces[tag] = node.get("typeface", "")
        missing = [t for t in ("latin", "ea", "cs") if t not in faces]
        if faces and missing:
            out.append(Finding(slide, "폰트", name,
                               f"latin·ea·cs 중 {', '.join(missing)} 누락 (현재: {faces})",
                               "세 개 모두 같은 typeface로 지정 — ea 누락 시 한글이 맑은 고딕으로 렌더"))
        for tag, face in faces.items():
            if face and not face.startswith("+") and not FONT_OK.match(face):
                out.append(Finding(slide, "폰트", name, f"{tag}='{face}' 는 Pretendard 계열이 아님",
                                   "Pretendard / Pretendard SemiBold / Pretendard Thin"))
        if rpr.get("b") == "1" and any("SemiBold" in f for f in faces.values()):
            out.append(Finding(slide, "폰트", name, "Pretendard SemiBold에 b=\"1\" 중복 적용",
                               "b 속성 제거 — 폰트 자체가 굵기를 갖는다"))


def check_size(el, slide, name, out, in_table):
    if in_table:
        return  # 표 셀은 9~10pt 예외
    for rpr in list(el.iter(f"{A}rPr")) + list(el.iter(f"{A}defRPr")):
        sz = rpr.get("sz")
        if sz and int(sz) / 100 < MIN_PT:
            out.append(Finding(slide, "크기", name, f"{int(sz)/100:g}pt — 하한 {MIN_PT:g}pt 미만",
                               "11pt 이상으로. 밀도가 문제면 페이지를 분할한다"))


def check_margin(slide, name, sh, out):
    if None in (sh.left, sh.top, sh.width, sh.height):
        return
    l, t, w, h = pt(sh.left), pt(sh.top), pt(sh.width), pt(sh.height)
    # 전폭 스트립(0~960)은 예외
    if l <= TOL and abs(l + w - SLIDE_W) <= TOL:
        return
    if l < MARGIN["left"] - TOL:
        out.append(Finding(slide, "마진", name, f"좌측 {l}pt < {MARGIN['left']:g}", "좌측 마진 46 기준으로 이동"))
    if l + w > MARGIN["right"] + TOL:
        out.append(Finding(slide, "마진", name, f"우측 끝 {round(l+w,2)}pt > {MARGIN['right']:g}", "폭을 줄이거나 좌측으로 이동"))
    if t < MARGIN["top"] - TOL:
        out.append(Finding(slide, "마진", name, f"상단 {t}pt < {MARGIN['top']:g}", "헤더 밴드 하단 +40 = y99 아래로"))
    if t + h > MARGIN["bottom"] + TOL:
        out.append(Finding(slide, "마진", name, f"하단 끝 {round(t+h,2)}pt > {MARGIN['bottom']:g}", "푸터 밴드 상단 −40 = y464 위로"))


def check_round(slide, name, sh, out):
    spPr = sh._element.find(f".//{A}prstGeom")
    if spPr is None or spPr.get("prst") != "roundRect":
        return
    # a:extLst의 ext가 아니라 spPr > a:xfrm > a:ext 를 읽어야 한다
    xfrm = sh._element.find(f".//{A}xfrm")
    ext = xfrm.find(f"{A}ext") if xfrm is not None else None
    if ext is None:
        return
    cx, cy = int(ext.get("cx", 0)), int(ext.get("cy", 0))
    if min(cx, cy) == 0:
        return
    want = min(round(ROUND_EMU / min(cx, cy) * 100000), 50000)
    gd = spPr.find(f".//{A}gd[@name='adj']")
    have = 16667  # roundRect 기본값
    if gd is not None:
        m = re.search(r"val\s+(-?\d+)", gd.get("fmla", ""))
        if m:
            have = int(m.group(1))
    if abs(have - want) > max(200, want * 0.02):
        out.append(Finding(slide, "라운드", name,
                           f"adj={have} (높이 {pt(cy)}pt 기준 실치수 4pt는 {want})",
                           f"adj를 {want}로 — adj는 도형마다 달라지는 게 정상이다"))


def check_subtraction(slide, name, sh, out):
    el = sh._element
    # spPr 은 p: 네임스페이스다. a: 로 찾으면 항상 None 이 나와 검사가 조용히 통과한다.
    spPr = el.find(f".//{P}spPr")
    if spPr is None:
        return None
    has_fill = spPr.find(f"{A}solidFill") is not None or spPr.find(f"{A}gradFill") is not None
    ln = spPr.find(f"{A}ln")
    has_line = ln is not None and ln.find(f"{A}noFill") is None and (
        ln.find(f"{A}solidFill") is not None or ln.get("w"))
    if has_fill and has_line:
        out.append(Finding(slide, "감산", name, "면 + 테두리 동시 사용",
                           "둘 중 하나만. 채움이 필요하면 #D5D2CA 면 하나만"))
    if el.find(f".//{A}outerShdw") is not None:
        out.append(Finding(slide, "감산", name, "그림자(outerShdw) 적용됨", "그림자는 전부 제거한다"))
    # lnRef(테마 테두리)가 있는데 a:ln override가 없으면 테마 선이 렌더된다.
    # 기본 생성 도형 대부분이 해당되므로 도형마다 찍으면 노이즈가 된다 → 슬라이드 단위로 집계한다.
    if el.find(f".//{A}lnRef") is not None and ln is None:
        return "lnRef"
    return None


def review(path, only=None):
    prs = Presentation(path)
    out = []
    lnref = defaultdict(int)
    for idx, slide in enumerate(prs.slides, start=1):
        if only and idx not in only:
            continue
        for name, sh in walk(slide.shapes):
            el = sh._element
            in_table = sh.has_table if hasattr(sh, "has_table") else False
            check_colors(el, idx, name, out)
            check_fonts(el, idx, name, out)
            check_size(el, idx, name, out, in_table)
            check_margin(idx, name, sh, out)
            check_round(idx, name, sh, out)
            if check_subtraction(idx, name, sh, out) == "lnRef":
                lnref[idx] += 1
        if lnref[idx]:
            out.append(Finding(idx, "감산", f"{lnref[idx]}개 도형",
                               "p:style의 lnRef만 있고 a:ln override가 없다 — 테마 테두리가 렌더된다",
                               "테두리를 지웠는데 안 사라지면 여기가 원인이다. <a:ln><a:noFill/></a:ln> 추가 또는 p:style 제거"))
    return out, len(prs.slides)


def main():
    ap = argparse.ArgumentParser(description="제안서 PPT 규격 검사 (ppt-design-system-v2 기준)")
    ap.add_argument("file")
    ap.add_argument("--slides", help="검사할 슬라이드 번호, 예: 29,36,59")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    only = {int(x) for x in a.slides.split(",")} if a.slides else None
    findings, total = review(a.file, only)

    if a.json:
        print(json.dumps(findings, ensure_ascii=False, indent=2))
        return 1 if findings else 0

    print(f"\n검사 파일: {a.file}  (전체 {total}장" + (f", 검사 {len(only)}장" if only else "") + ")")
    if not findings:
        print("\n규격 위반 없음.\n")
        return 0

    by_rule = defaultdict(int)
    for f in findings:
        by_rule[f["rule"]] += 1
    print("\n── 요약 ──")
    for rule in ("색상", "폰트", "크기", "마진", "라운드", "감산"):
        if by_rule[rule]:
            print(f"  {rule:<4} {by_rule[rule]:>3}건")
    print(f"  {'합계':<4} {len(findings):>3}건")

    by_slide = defaultdict(list)
    for f in findings:
        by_slide[f["slide"]].append(f)
    # 작업 순서: 마진 → 색·폰트 → 라운드 (배치가 바뀌면 나머지가 다시 틀어진다)
    order = {"마진": 0, "감산": 1, "색상": 2, "폰트": 3, "크기": 4, "라운드": 5}
    for slide in sorted(by_slide):
        print(f"\n── {slide}페이지 ──")
        for f in sorted(by_slide[slide], key=lambda x: order[x["rule"]]):
            print(f"  [{f['rule']}] {f['shape']}")
            print(f"        {f['detail']}")
            if f["fix"]:
                print(f"        {f['fix']}")
    print()
    return 1


if __name__ == "__main__":
    sys.exit(main())
